import docx

doc = docx.Document('c:/LIL/IDP_Report/IDP_EcoLifeBuddy_Report_Expanded_50Pages.docx')
breaks = 0
for p in doc.paragraphs:
    if 'w:type="page"' in p._p.xml:
        breaks += 1
print('Explicit Page Breaks:', breaks)
words = sum(len(p.text.split()) for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            words += len(cell.text.split())
print('Total Words:', words)
